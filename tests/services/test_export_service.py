"""Export (FR-601–FR-606): golden CSV files, XLSX read-back, FR-605, the PermissionError path.

The fixture database is fixed; the CSV goldens live in ``tests/fixtures`` and
are byte-compared, which catches accidental column, ordering or formatting
changes (PRD-02 §9). Regenerate deliberately with ``UPDATE_GOLDENS=1``.
"""

from __future__ import annotations

import os
import sqlite3
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

import pytest

from timetracker.core.clock import FakeClock
from timetracker.core.models import Dimension, NewEntry, RecordMethod
from timetracker.core.rounding import RoundingScope
from timetracker.data.entry_repo import EntryFilter, EntryRepo
from timetracker.data.label_repo import LabelRepo
from timetracker.services.export_service import (
    COLUMNS,
    ExportOptions,
    ExportService,
    build,
)

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
BRU = "Europe/Brussels"


def _seed(entries: EntryRepo, clients: LabelRepo, types: LabelRepo) -> None:
    """A Tuesday and a Wednesday of mixed work, incl. accents, quotes and an unlabelled row."""
    nike = clients.create("Nike")
    asics = clients.create("ASICS")
    email = types.create("Email/Chat")
    meeting = types.create("Meeting/Call")
    work = types.create("Work")
    tue = datetime(2026, 9, 8, 7, 0, tzinfo=UTC)  # 09:00 Brussels
    wed = tue + timedelta(days=1)
    rows = [
        (tue, 2, nike.id, email.id, "Réponse à Cédric", RecordMethod.QUICKADD),
        (tue + timedelta(minutes=10), 2, nike.id, email.id, None, RecordMethod.QUICKADD),
        (
            tue + timedelta(minutes=20),
            2,
            nike.id,
            email.id,
            'Said "ok"; moving on',
            RecordMethod.QUICKADD,
        ),
        (tue + timedelta(hours=1), 44, asics.id, meeting.id, "Kick-off", RecordMethod.STOPWATCH),
        (tue + timedelta(hours=3), 90, asics.id, work.id, "Drafting", RecordMethod.QUICKADD),
        (wed, 17, None, None, "unlabelled", RecordMethod.STOPWATCH),
        (wed + timedelta(hours=2), 26, nike.id, work.id, "multi\nline", RecordMethod.STOPWATCH),
    ]
    for start, minutes, client_id, type_id, note, method in rows:
        entries.insert(
            NewEntry(
                started_at_utc=start,
                ended_at_utc=start + timedelta(minutes=minutes),
                tz_name=BRU,
                duration_seconds=minutes * 60,
                record_method=method,
                client_id=client_id,
                type_id=type_id,
                note=note,
            )
        )


@pytest.fixture
def seeded(entries: EntryRepo, clients: LabelRepo, types: LabelRepo) -> EntryRepo:
    _seed(entries, clients, types)
    return entries


@pytest.fixture
def exporter(qapp, clock: FakeClock) -> ExportService:  # type: ignore[no-untyped-def]
    return ExportService(clock, None)


def _compare_or_update(actual: Path, golden: Path) -> None:
    if os.environ.get("UPDATE_GOLDENS"):
        golden.parent.mkdir(parents=True, exist_ok=True)
        golden.write_bytes(actual.read_bytes())
    assert golden.exists(), f"missing golden {golden.name}; run with UPDATE_GOLDENS=1"
    assert actual.read_bytes() == golden.read_bytes(), f"{actual.name} differs from {golden.name}"


@pytest.mark.parametrize(
    ("name", "minutes", "scope"),
    [
        ("none", 0, RoundingScope.PER_GROUP),
        ("6min_per_entry", 6, RoundingScope.PER_ENTRY),
        ("6min_per_group", 6, RoundingScope.PER_GROUP),
        ("15min_per_group", 15, RoundingScope.PER_GROUP),
    ],
)
def test_csv_golden(
    seeded: EntryRepo,
    exporter: ExportService,
    tmp_path: Path,
    name: str,
    minutes: int,
    scope: RoundingScope,
) -> None:
    out = tmp_path / f"{name}.csv"
    exporter.export_sync(
        seeded, ExportOptions(rounding_minutes=minutes, scope=scope, csv_delimiter=";"), out
    )
    raw = out.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    text = raw.decode("utf-8-sig")
    assert text.splitlines()[0] == ";".join(COLUMNS)
    assert "\r\n" in text  # RFC 4180 line endings from csv.writer(newline="")
    assert '"Said ""ok""; moving on"' in text  # quoting of the delimiter and quotes
    assert "Réponse à Cédric" in text
    _compare_or_update(out, FIXTURES / f"export_{name}.csv")


def test_csv_footer_records_rounding_and_filter(
    seeded: EntryRepo, exporter: ExportService, tmp_path: Path
) -> None:
    out = tmp_path / "x.csv"
    options = ExportOptions(
        flt=EntryFilter(date_from=date(2026, 9, 8), date_to=date(2026, 9, 8)),
        rounding_minutes=6,
        scope=RoundingScope.PER_GROUP,
        csv_delimiter=",",
    )
    exporter.export_sync(seeded, options, out)
    lines = out.read_text(encoding="utf-8-sig").splitlines()
    assert lines[0].startswith("Date,Start,End,")
    assert 'Rounding,"6 minutes, rounded up per day × client × type"' in lines
    assert "Filter,dates 2026-09-08 to 2026-09-08" in lines
    assert any(line.startswith("Total (h:mm),") for line in lines)
    assert sum(1 for line in lines[1:] if line and line[0].isdigit()) == 5  # Tuesday's five rows


def test_rounding_totals_per_scope(seeded: EntryRepo, clock: FakeClock) -> None:
    rows = seeded.export_rows(newest_first=False)
    now = clock.now_utc()
    none = build(rows, ExportOptions(rounding_minutes=0), now)
    entry6 = build(rows, ExportOptions(rounding_minutes=6, scope=RoundingScope.PER_ENTRY), now)
    group6 = build(rows, ExportOptions(rounding_minutes=6, scope=RoundingScope.PER_GROUP), now)
    assert none.raw_total_seconds == (2 + 2 + 2 + 44 + 90 + 17 + 26) * 60
    assert none.billed_total_seconds == none.raw_total_seconds
    # Per entry: 3×2→18, 44→48, 90→90, 17→18, 26→30 = 204 min
    assert entry6.billed_total_seconds == 204 * 60
    # Per group: the three Nike emails share a group → 6; 48 + 90 + 18 + 30 → 192 min
    assert group6.billed_total_seconds == 192 * 60
    assert [line.billed_seconds for line in group6.lines[:3]] == [120, 120, 120]


def test_xlsx_structure(seeded: EntryRepo, exporter: ExportService, tmp_path: Path) -> None:
    from openpyxl import load_workbook

    out = tmp_path / "week.xlsx"
    exporter.export_sync(
        seeded, ExportOptions(rounding_minutes=6, scope=RoundingScope.PER_GROUP), out
    )
    wb = load_workbook(out)
    assert wb.sheetnames == ["Time", "Export info"]
    ws = wb["Time"]
    assert [c.value for c in ws[1]] == list(COLUMNS)
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:I8"
    first = [c.value for c in ws[2]]
    assert first[:3] == ["2026-09-08", "09:00", "09:02"]
    assert isinstance(first[3], float) and first[3] == round(120 / 3600, 4)
    assert first[4] == timedelta(seconds=120)  # a real duration serial, read back as timedelta
    assert ws["E2"].number_format == "[h]:mm"
    assert ws["D2"].number_format == "0.00"
    assert ws["A9"].value == "Total"
    assert ws["D9"].value == "=SUBTOTAL(9,D2:D8)"
    assert ws["E9"].value == "=SUBTOTAL(9,E2:E8)"
    info = {row[0].value: row[1].value for row in wb["Export info"].iter_rows()}
    assert info["Rounding"] == "6 minutes, rounded up per day × client × type"
    assert info["Rounding scope"] == "per day × client × type"
    assert info["Entries"] == 7
    assert info["Raw total (h:mm)"] == "3:03"
    assert info["Billed total (h:mm)"] == "3:12"


def test_fr605_rounding_never_mutates_stored_data(
    seeded: EntryRepo, exporter: ExportService, tmp_path: Path, conn: sqlite3.Connection
) -> None:
    before = conn.execute(
        "SELECT id, duration_seconds, modified_at, is_edited FROM entry"
    ).fetchall()
    a = tmp_path / "a.csv"
    b = tmp_path / "b.csv"
    exporter.export_sync(seeded, ExportOptions(rounding_minutes=6, csv_delimiter=";"), a)
    exporter.export_sync(seeded, ExportOptions(rounding_minutes=15, csv_delimiter=";"), b)
    assert a.read_bytes() != b.read_bytes()  # different increment, different file
    after = conn.execute(
        "SELECT id, duration_seconds, modified_at, is_edited FROM entry"
    ).fetchall()
    assert [tuple(r) for r in before] == [tuple(r) for r in after]


def test_async_export_and_permission_error(qapp, qtbot, tmp_path: Path, clock: FakeClock) -> None:  # type: ignore[no-untyped-def]
    from timetracker.data.db import connect
    from timetracker.data.migrate import migrate

    db = tmp_path / "t.sqlite3"
    conn = connect(db)
    migrate(conn, db, clock=clock)
    clients = LabelRepo(conn, Dimension.CLIENT, clock)
    types = LabelRepo(conn, Dimension.TYPE, clock)
    _seed(EntryRepo(conn, clock), clients, types)

    svc = ExportService(clock, db)
    out = tmp_path / "async.xlsx"
    with qtbot.waitSignal(svc.export_finished, timeout=10_000) as blocker:
        svc.export(ExportOptions(), out)
    assert blocker.args == [out]
    assert out.exists()

    # A directory in the way is the same class of failure as "open in Excel".
    blocked = tmp_path / "blocked.csv"
    blocked.mkdir()
    with qtbot.waitSignal(svc.export_failed, timeout=10_000) as failed:
        svc.export(ExportOptions(), blocked)
    assert "blocked.csv" in failed.args[0] or "Export failed" in failed.args[0]
    conn.close()
