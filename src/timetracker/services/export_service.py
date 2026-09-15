"""CSV and XLSX export of the filtered log (FR-601–FR-606, PRD-02 §9).

Rounding is applied to the *rows being written* and recorded in the file
(FR-604); the database is never touched (FR-605). The two Duration columns
carry the billed values — that is what the recipient sums.

- **CSV**: ``utf-8-sig`` (the BOM is what makes Excel on Windows read UTF-8),
  RFC 4180 quoting via ``csv.writer(newline="")``, delimiter from settings
  (locale default), a footer stating the rounding and the filter.
- **XLSX**: ``openpyxl`` (lazy import). Decimal hours as a float, ``h:mm`` as an
  Excel duration serial with number format ``[h]:mm`` so it can be summed,
  header row frozen and auto-filtered, a totals row using ``SUBTOTAL(9, …)``
  so it respects filtering, and an ``Export info`` sheet with the metadata.

The file is built on a ``QThreadPool`` worker with its own read-only
connection (WAL makes readers concurrent with the writer, so the timer is
never blocked). ``build()`` and ``write_*()`` are also callable synchronously —
the golden-file tests use them directly.
"""

from __future__ import annotations

import csv
import sqlite3
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from PySide6.QtCore import QObject, QRunnable, QThreadPool, Signal

from timetracker import __version__
from timetracker.core.clock import Clock
from timetracker.core.duration import format_hm
from timetracker.core.rounding import BillableRow, RoundingScope, apply_rounding
from timetracker.core.timeutil import zone
from timetracker.data.db import connect
from timetracker.data.entry_repo import EntryFilter, EntryRepo, ExportRow

COLUMNS = (
    "Date",
    "Start",
    "End",
    "Duration (decimal hours)",
    "Duration (h:mm)",
    "Client",
    "Type",
    "Record method",
    "Note",
)


@dataclass(frozen=True, slots=True)
class ExportOptions:
    flt: EntryFilter = field(default_factory=EntryFilter)
    rounding_minutes: int = 0
    scope: RoundingScope = RoundingScope.PER_GROUP
    csv_delimiter: str = ";"
    columns: tuple[str, ...] = COLUMNS  # FR-603: selection is configurable


@dataclass(frozen=True, slots=True)
class ExportLine:
    date: str
    start: str
    end: str
    billed_seconds: int
    client: str
    type: str
    method: str
    note: str

    def cell(self, column: str) -> object:
        if column == "Date":
            return self.date
        if column == "Start":
            return self.start
        if column == "End":
            return self.end
        if column == "Duration (decimal hours)":
            return round(self.billed_seconds / 3600.0, 4)
        if column == "Duration (h:mm)":
            return format_hm(self.billed_seconds)
        if column == "Client":
            return self.client
        if column == "Type":
            return self.type
        if column == "Record method":
            return self.method
        if column == "Note":
            return self.note
        raise KeyError(column)


@dataclass(frozen=True, slots=True)
class ExportDocument:
    options: ExportOptions
    lines: list[ExportLine]
    generated_at: datetime
    raw_total_seconds: int
    billed_total_seconds: int

    @property
    def rounding_text(self) -> str:
        if self.options.rounding_minutes <= 0:
            return "none"
        return f"{self.options.rounding_minutes} minutes, rounded up {self.options.scope.display}"

    @property
    def filter_text(self) -> str:
        f = self.options.flt
        parts: list[str] = []
        if f.date_from or f.date_to:
            parts.append(f"dates {f.date_from or '…'} to {f.date_to or '…'}")
        if f.client_id is not None:
            parts.append(f"client id {f.client_id}")
        if f.type_id is not None:
            parts.append(f"type id {f.type_id}")
        if f.record_method is not None:
            parts.append(f"method {f.record_method.display}")
        if f.note_contains:
            parts.append(f"note contains {f.note_contains!r}")
        return "; ".join(parts) or "all entries"


def build(rows: list[ExportRow], options: ExportOptions, now: datetime) -> ExportDocument:
    """Pure: rows in, document out. Rounding never touches the rows (FR-605)."""
    billable = [
        BillableRow(
            r.entry.id,
            r.entry.local_date,
            r.entry.client_id,
            r.entry.type_id,
            r.entry.duration_seconds,
        )
        for r in rows
    ]
    billed = apply_rounding(billable, options.rounding_minutes, options.scope)
    lines: list[ExportLine] = []
    for r in rows:
        e = r.entry
        tz = zone(e.tz_name)
        lines.append(
            ExportLine(
                date=e.local_date.isoformat(),
                start=e.started_at_utc.astimezone(tz).strftime("%H:%M"),
                end=e.ended_at_utc.astimezone(tz).strftime("%H:%M"),
                billed_seconds=billed[e.id],
                client=r.client_name or "",
                type=r.type_name or "",
                method=e.record_method.display,
                note=e.note or "",
            )
        )
    return ExportDocument(
        options=options,
        lines=lines,
        generated_at=now,
        raw_total_seconds=sum(max(0, r.entry.duration_seconds) for r in rows),
        billed_total_seconds=sum(billed.values()),
    )


def write_csv(doc: ExportDocument, path: Path) -> None:
    cols = doc.options.columns
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=doc.options.csv_delimiter, quoting=csv.QUOTE_MINIMAL)
        w.writerow(cols)
        for line in doc.lines:
            w.writerow([_csv_cell(line.cell(c)) for c in cols])
        w.writerow([])
        w.writerow(["Total (h:mm)", format_hm(doc.billed_total_seconds)])
        w.writerow(
            ["Total (decimal hours)", _csv_cell(round(doc.billed_total_seconds / 3600.0, 4))]
        )
        w.writerow(["Rounding", doc.rounding_text])
        w.writerow(["Filter", doc.filter_text])
        w.writerow(["Exported", doc.generated_at.strftime("%Y-%m-%d %H:%M:%S %Z")])
        w.writerow(["Application", f"Time Tracker {__version__}"])


def _csv_cell(value: object) -> str:
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".") if value != int(value) else str(int(value))
    return str(value)


def write_xlsx(doc: ExportDocument, path: Path) -> None:
    from openpyxl import Workbook  # lazy: keeps cold start fast (R6)
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    cols = doc.options.columns
    wb = Workbook()
    ws = wb.active
    ws.title = "Time"
    ws.append(list(cols))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for line in doc.lines:
        row: list[object] = []
        for c in cols:
            if c == "Duration (h:mm)":
                row.append(line.billed_seconds / 86400.0)  # Excel duration serial
            else:
                row.append(line.cell(c))
        ws.append(row)

    n = len(doc.lines)
    first, last = 2, n + 1
    for idx, c in enumerate(cols, start=1):
        letter = get_column_letter(idx)
        if c == "Duration (decimal hours)":
            for r in range(first, last + 1):
                ws[f"{letter}{r}"].number_format = "0.00"
            ws[f"{letter}{last + 1}"] = f"=SUBTOTAL(9,{letter}{first}:{letter}{last})" if n else 0
            ws[f"{letter}{last + 1}"].number_format = "0.00"
            ws[f"{letter}{last + 1}"].font = Font(bold=True)
        elif c == "Duration (h:mm)":
            for r in range(first, last + 1):
                ws[f"{letter}{r}"].number_format = "[h]:mm"
            ws[f"{letter}{last + 1}"] = f"=SUBTOTAL(9,{letter}{first}:{letter}{last})" if n else 0
            ws[f"{letter}{last + 1}"].number_format = "[h]:mm"
            ws[f"{letter}{last + 1}"].font = Font(bold=True)
        width = max([len(c)] + [len(str(line.cell(c))) for line in doc.lines[:500]]) + 2
        ws.column_dimensions[letter].width = min(max(width, 8), 60)
    ws[f"A{last + 1}"] = "Total"
    ws[f"A{last + 1}"].font = Font(bold=True)
    ws.freeze_panes = "A2"
    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    for cell in ws[1]:
        cell.alignment = Alignment(horizontal="left")

    info = wb.create_sheet("Export info")
    info.append(["Exported", doc.generated_at.strftime("%Y-%m-%d %H:%M:%S %Z")])
    info.append(["Application", f"Time Tracker {__version__}"])
    info.append(["Filter", doc.filter_text])
    info.append(["Rounding", doc.rounding_text])
    info.append(["Rounding scope", doc.options.scope.display])
    info.append(["Entries", n])
    info.append(["Raw total (h:mm)", format_hm(doc.raw_total_seconds)])
    info.append(["Billed total (h:mm)", format_hm(doc.billed_total_seconds)])
    info.append(
        ["Note", "Durations are stored to the second; rounding was applied at export only."]
    )
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 70
    wb.save(path)


class ExportService(QObject):
    export_started = Signal()
    export_finished = Signal(object)  # Path
    export_failed = Signal(str)

    def __init__(self, clock: Clock, db_path: Path | None, parent: QObject | None = None) -> None:
        super().__init__(parent)
        self._clock = clock
        self._db_path = db_path  # None → in-memory tests use export_sync with a given repo
        self._pool = QThreadPool.globalInstance()

    # -- synchronous (tests, and the worker's body) ------------------------------

    def export_sync(self, repo: EntryRepo, options: ExportOptions, path: Path) -> Path:
        rows = repo.export_rows(options.flt, newest_first=False)
        doc = build(rows, options, self._clock.now_utc().astimezone(zone(self._clock.tz_name())))
        if path.suffix.lower() == ".xlsx":
            write_xlsx(doc, path)
        else:
            write_csv(doc, path)
        return path

    # -- asynchronous (the UI path) --------------------------------------------

    def export(self, options: ExportOptions, path: Path) -> None:
        """Run the export on a worker with its own read-only connection (§4.4, §9)."""
        if self._db_path is None:
            raise RuntimeError("ExportService needs a database path for asynchronous export")
        self.export_started.emit()
        self._pool.start(_ExportJob(self, self._db_path, options, path))


class _ExportJob(QRunnable):
    def __init__(
        self, service: ExportService, db_path: Path, options: ExportOptions, path: Path
    ) -> None:
        super().__init__()
        self._service = service
        self._db_path = db_path
        self._options = options
        self._path = path
        self.setAutoDelete(True)

    def run(self) -> None:
        conn: sqlite3.Connection | None = None
        try:
            conn = connect(self._db_path, read_only=True)
            repo = EntryRepo(conn, self._service._clock)  # noqa: SLF001 - same package
            self._service.export_sync(repo, self._options, self._path)
        except PermissionError:
            # Emitting from the worker is safe: receivers on the GUI thread get a
            # queued delivery automatically.
            self._service.export_failed.emit(
                f"{self._path.name} is open in another program (Excel?). "
                "Close it or choose another file name."
            )
            return
        except Exception as exc:  # noqa: BLE001 - reported to the UI, never swallowed
            self._service.export_failed.emit(f"Export failed: {exc}")
            return
        finally:
            if conn is not None:
                conn.close()
        self._service.export_finished.emit(self._path)
